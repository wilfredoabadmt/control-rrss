"""
Generador de Reportes Excel Reproducibles (.xlsx) — GAMEA Social Monitor
Principio XXIV: Reportes Reproducibles
Principio XXVI: Semántica de Indicadores
Principio XXVII: No Rankings
REQ-RPT-001, BR-RPT-001 a BR-RPT-005
"""

import hashlib
import io
import uuid
from datetime import UTC, datetime

import openpyxl
from core.audit.service import record_audit_event
from core.logging_config import get_correlation_id
from modules.iam.models import User
from modules.reporting.indicators import IndicatorEngine
from modules.reporting.models import ReportExecution
from modules.shared.enums import AuditAction
from modules.verification.models import Verification
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload


class ExcelReportGenerator:
    """
    Motor de compilación de reportes institucionales en formato OpenXML (.xlsx).
    Garantiza integridad con hash criptográfico SHA-256 y trazabilidad formal.
    """

    @staticmethod
    async def generate_verification_report(
        db: AsyncSession,
        current_user: User,
        campaign_title: str | None = None,
    ) -> tuple[bytes, str, ReportExecution]:
        """
        Genera el reporte oficial de verificaciones e indicadores con hash SHA-256.
        """
        cid = get_correlation_id()
        now = datetime.now(UTC)

        wb = openpyxl.Workbook()

        # Estilos institucionales GAMEA
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid") # Azul institucional
        thin_border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC"),
        )

        # ---------------------------------------------------------------------
        # Hoja 1: Metadatos y Ficha Técnica
        # ---------------------------------------------------------------------
        ws_meta = wb.active
        ws_meta.title = "Metadatos y Metodología"

        ws_meta["A1"] = "GOBIERNO AUTÓNOMO MUNICIPAL DE EL ALTO — GAMEA"
        ws_meta["A1"].font = Font(name="Calibri", size=14, bold=True, color="003366")
        ws_meta["A2"] = "PLATAFORMA DE MONITOREO Y AUDITORÍA DE REDES SOCIALES"
        ws_meta["A2"].font = Font(name="Calibri", size=11, bold=True, color="666666")

        coverage_sheet = await IndicatorEngine.calculate_observable_coverage_rate(db)
        verif_sheet = await IndicatorEngine.calculate_verification_rate(db)

        meta_rows = [
            ("Tipo de Reporte", "Estado de Verificación Epistémica y Cobertura Observable"),
            ("Versión del Generador", "1.0"),
            ("Fecha de Generación (UTC)", now.strftime("%Y-%m-%d %H:%M:%S UTC")),
            ("Usuario Solicitante", current_user.email),
            ("Correlation ID", cid),
            ("Campaña Seleccionada", campaign_title or "Consolidado Institucional General"),
            ("", ""),
            ("FICHA TÉCNICA: Tasa de Cobertura Observable", ""),
            ("Código de Indicador", coverage_sheet.code),
            ("Valor Obtenido", coverage_sheet.value_formatted),
            ("Numerador (Funcionarios Verificados)", str(coverage_sheet.numerator)),
            ("Denominador (Funcionarios Activos)", str(coverage_sheet.denominator)),
            ("Exclusiones Formales", coverage_sheet.exclusions),
            ("Tratamiento NOT_OBSERVABLE", coverage_sheet.not_observable_treatment),
            ("", ""),
            ("FICHA TÉCNICA: Tasa de Verificación Institucional", ""),
            ("Código de Indicador", verif_sheet.code),
            ("Valor Obtenido", verif_sheet.value_formatted),
            ("Numerador (Interacciones Confirmadas)", str(verif_sheet.numerator)),
            ("Denominador (Total Exigible)", str(verif_sheet.denominator)),
            ("Exclusiones Formales", verif_sheet.exclusions),
            ("Tratamiento NOT_OBSERVABLE", verif_sheet.not_observable_treatment),
            ("", ""),
            ("DECLARACIÓN CONSTITUCIONAL Y MARCO ÉTICO", ""),
            ("Principio V (No Inventar Datos)", "Las interacciones restringidas por plataformas se reportan como NOT_OBSERVABLE o API_RESTRICTED."),
            ("Principio XXVII (No Rankings)", "Este reporte institucional PROHÍBE taxativamente la generación de rankings o listas de mérito de funcionarios."),
        ]

        for r_idx, (k, v) in enumerate(meta_rows, start=4):
            ws_meta.cell(row=r_idx, column=1, value=k).font = Font(bold=True)
            ws_meta.cell(row=r_idx, column=2, value=v)

        ws_meta.column_dimensions["A"].width = 40
        ws_meta.column_dimensions["B"].width = 65

        # ---------------------------------------------------------------------
        # Hoja 2: Detalle de Verificaciones
        # ---------------------------------------------------------------------
        ws_data = wb.create_sheet(title="Verificaciones")

        headers = [
            "ID Verificación",
            "Funcionario ID",
            "Red Social",
            "Post ID",
            "Estado Epistémico",
            "Método de Verificación",
            "Fecha Verificación (UTC)",
            "Explicación Comprensible",
        ]

        for col_idx, h in enumerate(headers, start=1):
            cell = ws_data.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Consultar verificaciones registradas
        stmt_verifs = (
            select(Verification)
            .options(
                selectinload(Verification.interaction),
            )
            .order_by(Verification.verified_at.desc())
            .limit(1000)
        )
        verifs: list[Verification] = list((await db.execute(stmt_verifs)).scalars().all())

        for row_idx, v in enumerate(verifs, start=2):
            ws_data.cell(row=row_idx, column=1, value=str(v.id))
            ws_data.cell(row=row_idx, column=2, value=v.employee_id or "NO_CRUZADO")
            ws_data.cell(row=row_idx, column=3, value=v.interaction.source_platform or "Social Platform" if v.interaction else "N/A")
            ws_data.cell(row=row_idx, column=4, value=v.interaction.external_post_id if v.interaction else "N/A")
            ws_data.cell(row=row_idx, column=5, value=v.verification_status)
            ws_data.cell(row=row_idx, column=6, value=v.verification_method)
            ws_data.cell(row=row_idx, column=7, value=v.verified_at.strftime("%Y-%m-%d %H:%M UTC"))
            ws_data.cell(row=row_idx, column=8, value=v.explanation)

            for c in range(1, 9):
                ws_data.cell(row=row_idx, column=c).border = thin_border

        for col_letter in ["A", "B", "C", "D", "E", "F", "G", "H"]:
            ws_data.column_dimensions[col_letter].width = 25
        ws_data.column_dimensions["H"].width = 50

        # Guardar en memoria y calcular hash SHA-256 (Principio XXIV)
        buffer = io.BytesIO()
        wb.save(buffer)
        excel_bytes = buffer.getvalue()
        file_hash = hashlib.sha256(excel_bytes).hexdigest()

        # Registrar ReportExecution inmutable
        execution = ReportExecution(
            id=uuid.uuid4(),
            report_type="VERIFICATION_STATUS",
            report_version="1.0",
            requested_by_user_id=current_user.id,
            generated_at=now,
            parameters={"campaign": campaign_title, "max_rows": len(verifs)},
            file_hash=file_hash,
            file_path=f"/reports/verification_{now.strftime('%Y%m%d_%H%M%S')}.xlsx",
            row_count=len(verifs),
            status="COMPLETED",
        )
        db.add(execution)

        await record_audit_event(
            db=db,
            action=AuditAction.EXPORT,
            entity_name="ReportExecution",
            entity_id=str(execution.id),
            user_id=str(current_user.id),
            user_email=current_user.email,
            new_state={"file_hash": file_hash, "row_count": len(verifs)},
            details={"report_type": "VERIFICATION_STATUS", "operation": "EXCEL_GENERATION"},
            correlation_id=cid,
        )

        await db.refresh(execution)

        return excel_bytes, file_hash, execution
