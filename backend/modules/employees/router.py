"""
Routers para Directorio de Funcionarios, Estructura Organizacional y Cargos — GAMEA Social Monitor
"""

import pathlib
import re
import uuid

from core.pagination import PageResponse
from core.security.auth import get_current_user
from core.security.encryption import decrypt_field
from core.security.rbac import require_roles
from database import get_async_db
from dependencies import PaginationDep
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import FileResponse
from modules.employees.importer import EmployeePayrollImporter
from modules.employees.models import Employee, OrganizationalUnit, Position
from modules.employees.schemas import (
    EmployeeCreate,
    EmployeeBulkDeleteRequest,
    EmployeeBulkDeleteResponse,
    EmployeeDetailResponse,
    EmployeeHistoryResponse,
    EmployeeImportReport,
    EmployeeResponse,
    EmployeeUpdate,
    OrganizationalUnitCreate,
    OrganizationalUnitResponse,
    OrganizationalUnitTreeNode,
    OrganizationalUnitUpdate,
    PositionCreate,
    PositionResponse,
)
from modules.employees.service import EmployeeService
from modules.iam.models import User
from modules.shared.enums import EmployeeStatus, UserRole
from modules.shared.exceptions import EntityNotFoundException, ValidationException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

org_units_router = APIRouter()
positions_router = APIRouter()
employees_router = APIRouter()


# Roles con permiso de gestión de funcionarios, unidades y nómina
EMPLOYEE_MANAGE_ROLES = (
    UserRole.SUPER_ADMIN,
    UserRole.DIRECTOR,
    UserRole.COMMUNICATIONS_LEAD,
    UserRole.ANALYST,
    UserRole.OPERATOR,
)

# -----------------------------------------------------------------------------
# Endpoints de Unidades Organizacionales (/api/v1/org-units)
# -----------------------------------------------------------------------------

@org_units_router.get("/tree", response_model=list[OrganizationalUnitTreeNode])
async def get_org_units_tree(
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(get_current_user),
):
    """Retorna la jerarquía completa en árbol de las unidades del GAMEA."""
    return await EmployeeService.get_org_units_tree(db)


@org_units_router.get("/", response_model=list[OrganizationalUnitResponse])
async def list_org_units(
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(get_current_user),
):
    """Retorna la lista plana de unidades organizacionales."""
    stmt = select(OrganizationalUnit).order_by(OrganizationalUnit.name)
    units = list((await db.execute(stmt)).scalars().all())
    return [OrganizationalUnitResponse.model_validate(u) for u in units]


@org_units_router.post("/", response_model=OrganizationalUnitResponse, status_code=status.HTTP_201_CREATED)
async def create_org_unit(
    org_in: OrganizationalUnitCreate,
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(require_roles(*EMPLOYEE_MANAGE_ROLES)),
):
    """Crea una nueva unidad organizacional."""
    try:
        unit = await EmployeeService.create_org_unit(db, org_in, current_user)
        return OrganizationalUnitResponse.model_validate(unit)
    except ValidationException as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=e.message) from e


@org_units_router.patch("/{unit_id}", response_model=OrganizationalUnitResponse)
async def update_org_unit(
    unit_id: uuid.UUID,
    org_in: OrganizationalUnitUpdate,
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(require_roles(*EMPLOYEE_MANAGE_ROLES)),
):
    """Actualiza una unidad organizacional existente."""
    try:
        unit = await EmployeeService.update_org_unit(db, unit_id, org_in, current_user)
        return OrganizationalUnitResponse.model_validate(unit)
    except EntityNotFoundException as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=e.message) from e
    except ValidationException as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=e.message) from e


# -----------------------------------------------------------------------------
# Endpoints de Cargos (/api/v1/positions)
# -----------------------------------------------------------------------------

@positions_router.get("/", response_model=list[PositionResponse])
async def list_positions(
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(get_current_user),
):
    """Retorna el catálogo institucional de cargos."""
    stmt = select(Position).order_by(Position.title)
    positions = list((await db.execute(stmt)).scalars().all())
    return [PositionResponse.model_validate(p) for p in positions]


@positions_router.post("/", response_model=PositionResponse, status_code=status.HTTP_201_CREATED)
async def create_position(
    pos_in: PositionCreate,
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(require_roles(*EMPLOYEE_MANAGE_ROLES)),
):
    """Crea un nuevo cargo institucional."""
    pos = await EmployeeService.create_position(db, pos_in, current_user)
    return PositionResponse.model_validate(pos)


# -----------------------------------------------------------------------------
# Endpoints de Funcionarios (/api/v1/employees)
# -----------------------------------------------------------------------------

def _format_employee_response(emp: Employee, current_user: User) -> EmployeeResponse:
    """Aplica protección PII: Oculta CI si el usuario no tiene permisos suficientes."""
    user_roles = {r.name for r in current_user.roles}
    privileged_roles = {
        UserRole.SUPER_ADMIN.value,
        UserRole.AUDITOR.value,
        UserRole.DIRECTOR.value,
        UserRole.COMMUNICATIONS_LEAD.value,
        UserRole.ANALYST.value,
        UserRole.OPERATOR.value,
    }

    # Si es privilegiado, descifra la cédula; si es visor, la oculta/enmascara
    doc_number = None
    if user_roles.intersection(privileged_roles):
        doc_number = decrypt_field(emp.document_number_encrypted)
    else:
        doc_number = "••••••"

    org_unit_name = None
    parent_unit_name = None
    try:
        if getattr(emp, "direction_name", None):
            parent_unit_name = emp.direction_name
        if getattr(emp, "organizational_unit", None):
            org_unit_name = emp.organizational_unit.name
            if getattr(emp.organizational_unit, "parent", None):
                parent_unit_name = emp.organizational_unit.parent.name
    except Exception:
        pass

    if parent_unit_name:
        parent_unit_name = re.sub(r'(?i)alcaldesa', 'Alcalde', parent_unit_name)
    if org_unit_name:
        org_unit_name = re.sub(r'(?i)alcaldesa', 'Alcalde', org_unit_name)

    position_title = None
    try:
        if getattr(emp, "position", None):
            position_title = emp.position.title
    except Exception:
        pass

    # Redes sociales asociadas
    fb_acc = None
    tt_acc = None
    try:
        socials = getattr(emp, "social_accounts", None)
        if socials:
            for sa in socials:
                p_code = getattr(getattr(sa, "platform", None), "name", "") or getattr(getattr(sa, "platform", None), "code", "")
                url = (getattr(sa, "profile_url", None) or "").lower()
                username = getattr(sa, "current_username", None) or ""
                if p_code == "FACEBOOK" or "facebook" in url:
                    fb_acc = username or getattr(sa, "profile_url", None)
                elif p_code == "TIKTOK" or "tiktok" in url:
                    tt_acc = username or getattr(sa, "profile_url", None)
    except Exception:
        pass

    return EmployeeResponse(
        id=emp.employee_id,
        employee_id=emp.employee_id,
        first_name=emp.first_name,
        last_name=emp.last_name,
        document_number=doc_number,
        id_document=doc_number,
        email=getattr(emp, "email", None) or f"{emp.first_name.lower().split()[0]}.{emp.last_name.lower().split()[0]}@elalto.gob.bo",
        organizational_unit_id=emp.organizational_unit_id,
        org_unit_name=org_unit_name,
        parent_unit_name=parent_unit_name,
        position_id=emp.position_id,
        position_title=position_title,
        facebook_account=fb_acc,
        tiktok_account=tt_acc,
        status=emp.status,
        is_active=emp.status == EmployeeStatus.ACTIVE.value,
        hire_date=emp.hire_date,
        termination_date=emp.termination_date,
        created_at=emp.created_at,
        updated_at=emp.updated_at,
    )


@employees_router.get("/", response_model=PageResponse[EmployeeResponse])
async def list_employees(
    pagination: PaginationDep,
    unit_id: uuid.UUID | None = Query(None, description="Filtrar por unidad organizacional"),
    direction: str | None = Query(None, description="Filtrar por dirección superior"),
    status: str | None = Query(None, description="Filtrar por estado"),
    search: str | None = Query(None, description="Búsqueda por nombre o ID"),
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(get_current_user),
):
    """Consulta paginada del directorio de funcionarios con aislamiento individual."""
    employees, total = await EmployeeService.list_employees(
        db=db,
        unit_id=unit_id,
        direction=direction,
        status=status,
        search=search,
        offset=pagination.offset,
        limit=pagination.limit,
        current_user=current_user,
    )
    items = [_format_employee_response(e, current_user) for e in employees]
    return PageResponse.create(items=items, total=total, page=pagination.page, page_size=pagination.page_size)


@employees_router.post("/", response_model=EmployeeResponse, status_code=status.HTTP_201_CREATED)
async def create_employee(
    emp_in: EmployeeCreate,
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(require_roles(*EMPLOYEE_MANAGE_ROLES)),
):
    """Crea un funcionario institucional."""
    try:
        emp = await EmployeeService.create_employee(db, emp_in, current_user)
        return _format_employee_response(emp, current_user)
    except ValidationException as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=e.message) from e


@employees_router.get("/import/template")
async def download_import_template(
    file_format: str = Query("xlsx", alias="format", description="Formato de plantilla: 'xlsx' o 'csv'"),
    current_user: User = Depends(get_current_user),
):
    """
    Descarga la plantilla oficial (.xlsx o .csv) para importar nómina de funcionarios.
    Fila 1: nombres, apellidos, unidad, direccion, cuenta_facebook, cuenta_tiktok.
    Sin datos de ejemplo, limpia para llenado institucional.
    """
    is_csv = file_format.lower() == "csv"
    ext = "csv" if is_csv else "xlsx"
    target_filename = f"plantilla_funcionarios_gamea.{ext}"
    media_type = "text/csv" if is_csv else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    candidates = [
        pathlib.Path(__file__).resolve().parents[3] / "extras" / f"plantilla_funcionarios.{ext}",
        pathlib.Path(__file__).resolve().parents[2] / "extras" / f"plantilla_funcionarios.{ext}",
        pathlib.Path(f"/app/extras/plantilla_funcionarios.{ext}"),
        pathlib.Path(f"extras/plantilla_funcionarios.{ext}"),
    ]
    template_path = next((p for p in candidates if p.is_file()), None)

    if template_path:
        return FileResponse(
            path=str(template_path),
            filename=target_filename,
            media_type=media_type,
            headers={"Content-Disposition": f"attachment; filename={target_filename}"},
        )

    # Si no existe en disco, generar al vuelo
    if not is_csv:
        import io
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from fastapi.responses import Response

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Funcionarios GAMEA"
        headers = ["nombres", "apellidos", "unidad", "direccion", "cuenta_facebook", "cuenta_tiktok"]
        ws.append(headers)
        ws.row_dimensions[1].height = 28

        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="334155"),
            right=Side(style="thin", color="334155"),
            top=Side(style="thin", color="334155"),
            bottom=Side(style="thin", color="334155"),
        )
        col_widths = {"A": 22, "B": 24, "C": 36, "D": 36, "E": 30, "F": 22}

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = col_widths.get(col_letter, 20)

        ws.auto_filter.ref = "A1:F1"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return Response(
            content=buf.getvalue(),
            media_type=media_type,
            headers={"Content-Disposition": f"attachment; filename={target_filename}"},
        )

    from fastapi.responses import PlainTextResponse
    csv_content = "\ufeffnombres,apellidos,unidad,direccion,cuenta_facebook,cuenta_tiktok\n , , , , , \n"
    return PlainTextResponse(
        content=csv_content,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={target_filename}"},
    )


@employees_router.post("/import", response_model=EmployeeImportReport)
async def import_employees_payroll(
    file: UploadFile = File(..., description="Archivo Excel (.xlsx) o CSV de nómina"),
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(require_roles(*EMPLOYEE_MANAGE_ROLES)),
):
    """
    Importación y sincronización de nómina institucional (REQ-EMP-003).
    Idempotente: Detecta altas, bajas y transferencias organizacionales.
    """
    if not file.filename:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Nombre de archivo inválido.")

    contents = await file.read()
    report = await EmployeePayrollImporter.import_payroll_file(
        db=db,
        file_content=contents,
        filename=file.filename,
        current_user=current_user,
    )
    return report


@employees_router.get("/{employee_id}", response_model=EmployeeDetailResponse)
async def get_employee(
    employee_id: str,
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(get_current_user),
):
    """Consulta la ficha completa de un funcionario."""
    emp = await EmployeeService.get_employee_by_id(db, employee_id)
    if not emp or not EmployeeService.is_accessible_by_user(emp, current_user):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Funcionario no encontrado.")

    base_resp = _format_employee_response(emp, current_user)
    history_items = [EmployeeHistoryResponse.model_validate(h) for h in emp.history]

    return EmployeeDetailResponse(
        **base_resp.model_dump(),
        organizational_unit=OrganizationalUnitResponse.model_validate(emp.organizational_unit) if emp.organizational_unit else None,
        position=PositionResponse.model_validate(emp.position) if emp.position else None,
        history=history_items,
    )


@employees_router.get("/{employee_id}/history", response_model=list[EmployeeHistoryResponse])
async def get_employee_history(
    employee_id: str,
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(get_current_user),
):
    """
    Retorna la línea temporal de transferencias y cambios organizacionales del funcionario (T-204).
    """
    emp = await EmployeeService.get_employee_by_id(db, employee_id)
    if not emp or not EmployeeService.is_accessible_by_user(emp, current_user):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Funcionario no encontrado.")
    return [EmployeeHistoryResponse.model_validate(h) for h in emp.history]


@employees_router.patch("/{employee_id}", response_model=EmployeeResponse)
async def update_employee(
    employee_id: str,
    emp_in: EmployeeUpdate,
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(require_roles(*EMPLOYEE_MANAGE_ROLES)),
):
    """
    Actualiza datos de un funcionario. Si cambia de unidad o cargo, registra en el historial (T-204).
    """
    try:
        emp = await EmployeeService.update_employee(db, employee_id, emp_in, current_user)
        return _format_employee_response(emp, current_user)
    except EntityNotFoundException as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=e.message) from e
    except ValidationException as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=e.message) from e


@employees_router.post("/bulk-delete", response_model=EmployeeBulkDeleteResponse, status_code=status.HTTP_200_OK)
async def bulk_delete_employees(
    payload: EmployeeBulkDeleteRequest,
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(require_roles(*EMPLOYEE_MANAGE_ROLES)),
):
    """
    Eliminación masiva / en lote de funcionarios (física o baja lógica).
    Permite procesar múltiples funcionarios simultáneamente de forma atómica y segura.
    """
    res = await EmployeeService.bulk_delete_employees(
        db=db,
        employee_ids=payload.employee_ids,
        current_user=current_user,
        permanent=payload.permanent,
        reason=payload.reason,
    )
    return EmployeeBulkDeleteResponse(**res)


@employees_router.delete("/{employee_id}", status_code=status.HTTP_200_OK)
async def delete_employee(
    employee_id: str,
    reason: str | None = Query("Desvinculación institucional", description="Motivo de la baja"),
    permanent: bool = Query(False, description="Eliminación física definitiva de la base de datos"),
    db: AsyncSession = Depends(get_async_db),
    current_user: User = Depends(require_roles(*EMPLOYEE_MANAGE_ROLES)),
):
    """
    Baja lógica o eliminación de funcionario (BR-EMP-004).
    """
    try:
        await EmployeeService.delete_employee(db, employee_id, current_user, permanent=permanent, reason=reason)
        msg = "Funcionario eliminado permanentemente correctamente." if permanent else "Funcionario dado de baja lógica correctamente."
        return {"detail": msg}
    except EntityNotFoundException as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=e.message) from e
