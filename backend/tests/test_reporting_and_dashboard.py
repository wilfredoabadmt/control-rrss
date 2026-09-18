"""
Pruebas de Reportes Excel Reproducibles y Dashboards — GAMEA Social Monitor
Fase 6: Reportes y Dashboard
Principio XXIV: Reportes Reproducibles
Principio XXV: Dashboard Verificable
Principio XXVI: Semántica de Indicadores
Principio XXVII: No Rankings
REQ-RPT-001, REQ-RPT-003, REQ-DSH-001, REQ-DSH-002
"""

import io
import uuid

import openpyxl
import pytest
import pytest_asyncio
from config import settings
from database import Base, get_async_db
from fastapi.testclient import TestClient
from main import app
from modules.employees.models import Employee
from modules.iam.models import Role, User
from modules.interactions.models import Interaction
from modules.publications.models import Publication
from modules.reporting.generator import ExcelReportGenerator
from modules.reporting.indicators import IndicatorEngine
from modules.shared.enums import (
    SocialPlatformType,
    UserRole,
    VerificationStatus,
)
from modules.social_accounts.models import SocialPlatform
from modules.social_accounts.seed import seed_social_platforms
from modules.verification.models import Verification
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine


@pytest_asyncio.fixture
async def async_db():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async_session = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    async with async_session() as session:
        await seed_social_platforms(session)

        # Crear funcionario
        emp = Employee(
            employee_id="FUNC-2026-999",
            first_name="Guillermo",
            last_name="Chavez",
            document_number_encrypted="enc_doc",
            document_hash="hash_doc",
            status="ACTIVE",
        )
        session.add(emp)

        # Crear plataforma y publicación
        plat = (await session.execute(
            SocialPlatform.__table__.select().where(SocialPlatform.name == SocialPlatformType.FACEBOOK.value)
        )).first()

        pub = Publication(
            id=uuid.uuid4(),
            platform_id=plat.id,
            external_post_id="post_report_test_01",
            content_text="Post para prueba de reportes",
            is_monitored=True,
        )
        session.add(pub)

        # Crear interacción
        inter = Interaction(
            id=uuid.uuid4(),
            publication_id=pub.id,
            platform_id=plat.id,
            interaction_type="COMMENT",
            external_interaction_id="comm_rep_01",
            external_author_id="user_rep_01",
            content_text="Comentario de prueba",
            capture_method="POLLING",
            data_origin_type="EMPLOYEE_INTERACTION_OFFICIAL",
        )
        session.add(inter)

        # Crear verificación CONFIRMED
        verif = Verification(
            id=uuid.uuid4(),
            interaction_id=inter.id,
            employee_id="FUNC-2026-999",
            verification_status=VerificationStatus.CONFIRMED.value,
            verification_method="AUTOMATIC_CROSS_REFERENCE",
            explanation="Verificación de prueba para reporte",
        )
        session.add(verif)

        await session.commit()
        yield session

    await engine.dispose()


@pytest.fixture
def mock_admin_user():
    admin_role = Role(
        id=uuid.uuid4(),
        name=UserRole.SUPER_ADMIN.value,
        description="Super Administrador",
    )
    u = User(
        id=uuid.uuid4(),
        email="admin.reportes@elalto.gob.bo",
        full_name="Admin Reportes",
        password_hash="hash",
    )
    u.roles.append(admin_role)
    return u


# -----------------------------------------------------------------------------
# T-601: Fichas Técnicas e Indicadores (REQ-RPT-003 / Principio XXVI)
# -----------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_indicator_engine_calculations(async_db: AsyncSession):
    """
    REQ-RPT-003: Cálculo de Tasa de Cobertura Observable y Tasa de Verificación con Ficha Técnica.
    """
    coverage = await IndicatorEngine.calculate_observable_coverage_rate(async_db)
    assert coverage.code == "IND-COV-001"
    assert coverage.numerator == 1.0
    assert coverage.denominator == 1.0
    assert coverage.value == 100.0
    assert "Excluye funcionarios" in coverage.exclusions
    assert "NOT_OBSERVABLE" in coverage.not_observable_treatment

    verification_rate = await IndicatorEngine.calculate_verification_rate(async_db)
    assert verification_rate.code == "IND-VER-002"
    assert verification_rate.numerator == 1.0
    assert verification_rate.denominator == 1.0
    assert verification_rate.value == 100.0

    status_dist = await IndicatorEngine.get_verification_status_distribution(async_db)
    assert status_dist[VerificationStatus.CONFIRMED.value] == 1


# -----------------------------------------------------------------------------
# T-600: Generación de Reportes Excel con Integridad Criptográfica (REQ-RPT-001)
# -----------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_excel_report_generation_reproducibility(async_db: AsyncSession, mock_admin_user: User):
    """
    REQ-RPT-001 & Principio XXIV: Generación de .xlsx reproducible con metadatos y hash SHA-256.
    """
    excel_bytes, file_hash, execution = await ExcelReportGenerator.generate_verification_report(
        db=async_db,
        current_user=mock_admin_user,
        campaign_title="Campaña Reportes 2026",
    )

    assert len(excel_bytes) > 0
    assert len(file_hash) == 64  # SHA-256 hexadecimal
    assert execution.file_hash == file_hash
    assert execution.row_count == 1
    assert execution.status == "COMPLETED"

    # Verificar que openpyxl puede parsear el binario generado
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Metadatos y Metodología" in wb.sheetnames
    assert "Verificaciones" in wb.sheetnames

    # Verificar ficha técnica en la hoja de metadatos
    ws_meta = wb["Metadatos y Metodología"]
    assert "GOBIERNO AUTÓNOMO MUNICIPAL DE EL ALTO" in ws_meta["A1"].value
    assert "Principio XXVII (No Rankings)" in [ws_meta.cell(row=r, column=1).value for r in range(1, 40)]


# -----------------------------------------------------------------------------
# T-602 & T-603: Endpoints de Dashboard Operativo y Ejecutivo (REQ-DSH-001 / 002)
# -----------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_dashboard_api_endpoints(async_db: AsyncSession, mock_admin_user: User):
    """
    REQ-DSH-001 / REQ-DSH-002: Consulta de dashboard operativo y ejecutivo con datos del backend.
    """
    app.dependency_overrides[get_async_db] = lambda: async_db
    from core.security.auth import get_current_user
    app.dependency_overrides[get_current_user] = lambda: mock_admin_user

    client = TestClient(app)

    # 1. Dashboard Operativo
    resp_op = client.get(f"{settings.API_V1_STR}/dashboard/operational")
    assert resp_op.status_code == 200
    data_op = resp_op.json()
    assert len(data_op["platforms"]) >= 2
    assert data_op["monitored_publications_count"] >= 1
    assert data_op["total_interactions_count"] >= 1

    # 2. Dashboard Ejecutivo
    resp_exec = client.get(f"{settings.API_V1_STR}/dashboard/executive")
    assert resp_exec.status_code == 200
    data_exec = resp_exec.json()
    assert data_exec["observable_coverage_rate"]["code"] == "IND-COV-001"
    assert data_exec["verification_rate"]["code"] == "IND-VER-002"
    assert "prohíbe taxativamente la elaboración de rankings" in data_exec["constitutional_disclaimer"]

    # 3. Exportación Excel vía HTTP
    resp_export = client.post(
        f"{settings.API_V1_STR}/reports/export-excel",
        json={"report_type": "VERIFICATION_STATUS", "campaign_title": "Campaña Test"},
    )
    assert resp_export.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp_export.headers["content-type"]
    assert "X-Report-SHA256" in resp_export.headers

    app.dependency_overrides.clear()
